#!/Users/calindgr/.pyenv/versions/3.12.4/bin/python
"""
__author__ = "Chris Lindgren"
__email__ = "calindgren@ncsu.edu"
Parse text with spaCy and print the output in CoNLL-U format.
References:
https://spacy.io/
http://universaldependencies.org/format.html
"""
import argparse
from contextlib import redirect_stdout
from pathlib import Path
import re
import shutil
from docx import Document
from syntaxcomp.complexity import TextComplexity,SentenceComplexity
from conllu import parse
import spacy

def segment_txt_as_docx(Path__input_folder, Path__output_folder, unit_type, output_format="docx"):
  """Convert CoNLL-U text files into files containing selected segments.

  Each source text file is parsed for sentence, T-unit, and clause structure.
  The requested segments are written one per paragraph (DOCX) or line (TXT),
  retaining the input file's relative directory and filename. When
  ``unit_type`` is ``"all"``, separate trees are created for sentences,
  T-units, and clauses beneath the output folder.

  Args:
      Path__input_folder: Root directory containing CoNLL-U ``.txt`` files.
      Path__output_folder: Root directory for generated segmented files.
      unit_type: One of ``"s"``, ``"t"``, ``"c"``, or ``"all"``.
      output_format: Either ``"docx"`` or ``"txt"``.

  Returns:
      None. Segmented files in the requested format are written to disk.
  """
  CONLLU_folder_base = Path__input_folder  # CONLLU folder base
  SEGMENTED_folder_base = Path__output_folder  # segmented folder base
  unit_types = ("s", "t", "c") if unit_type == "all" else (unit_type,)
  unit_foldernames = {
    "s": "sentences",
    "t": "t_units",
    "c": "clauses",
  }

  for path_conllu in sorted(CONLLU_folder_base.rglob("*.txt")):
    relative_path = path_conllu.relative_to(CONLLU_folder_base)

    with path_conllu.open("r", encoding="utf-8") as text_file:
      str__conllu = text_file.read()

    anno = parse(str__conllu)
    tc = TextComplexity(anno)
    segmented_docs = {unit: [] for unit in unit_types}

    for sentence in tc.sentences:
      sc = SentenceComplexity(sentence)

      if "s" in segmented_docs:
        re__sentence_pattern = re.compile(r"(\stext\:\s\")(.{1,})(\"\}\>)", flags=re.MULTILINE)
        matches = re__sentence_pattern.finditer(str(sentence))
        if matches:
          for match in matches:
            segmented_docs["s"].append(match.group(2))
        else:
          print('No sentences found')

      if "t" in segmented_docs:
        for tu in sc.t_units:
          if 'text' in tu:
            segmented_docs["t"].append(tu['text'])

      if "c" in segmented_docs:
        for clause in sc.clauses:
          if 'text' in clause:
            segmented_docs["c"].append(clause['text'])

    for unit, segments in segmented_docs.items():
      output_folder = SEGMENTED_folder_base
      if unit_type == "all":
        output_folder /= unit_foldernames[unit]
      path_segmented_file = output_folder / relative_path.with_suffix(f".{output_format}")
      path_segmented_file.parent.mkdir(parents=True, exist_ok=True)

      if output_format == "docx":
        segmented_docx = Document()
        for segment in segments:
          segmented_docx.add_paragraph(segment)
        segmented_docx.save(path_segmented_file)
      else:
        with path_segmented_file.open("w", encoding="utf-8") as output_file:
          for segment in segments:
            output_file.write(f"{segment}\n")

def convert_docx_to_txt(Path__input_folder, Path__output_folder):
  """Extract paragraph text from every DOCX file under an input directory.

  The source directory is searched recursively. Relative paths are preserved
  under the output directory, and temporary Microsoft Word lock files are
  skipped.

  Args:
      Path__input_folder: Root directory containing source ``.docx`` files.
      Path__output_folder: Root directory for the generated ``.txt`` files.

  Returns:
      None. UTF-8 text files are written to disk.
  """
  DOCX_folder_base = Path__input_folder  # DOCX folder base
  TXT_folder_base = Path__output_folder  # TXT folder base

  for path_docx in sorted(DOCX_folder_base.rglob("*.docx")):
    if path_docx.name.startswith("~$"):
      continue

    relative_path = path_docx.relative_to(DOCX_folder_base)
    path_txt = TXT_folder_base / relative_path.with_suffix(".txt")
    path_txt.parent.mkdir(parents=True, exist_ok=True)

    document = Document(path_docx)
    with path_txt.open('w', encoding='utf-8') as text_file:
      for paragraph in document.paragraphs:
        text_file.write(paragraph.text)
        text_file.write('\n')

def copy_conllu_files(input_folder, output_folder):
  """Copy CoNLL-U text files to a destination while preserving relative paths.

  Args:
      input_folder: Root directory containing generated CoNLL-U ``.txt`` files.
      output_folder: Destination root for the copied files.

  Returns:
      None. Files are copied with their metadata preserved.
  """
  conllu_root = Path(input_folder)
  destination_root = Path(output_folder)
  for path_conllu in sorted(conllu_root.rglob("*.txt")):
    relative_path = path_conllu.relative_to(conllu_root)
    destination = destination_root / relative_path
    destination.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(path_conllu, destination)

def sentences_to_conllu(doc, sent_id, prefix = ""):
  """Print sentences from a spaCy document in CoNLL-U format.

  Writes a sentence header, its text, and a CoNLL-U token row for each token
  to standard output. Token rows include IDs, forms, lemmas, part-of-speech
  tags, morphology, dependency heads, and dependency relations.

  Args:
      doc: A spaCy ``Doc`` whose sentences and dependency annotations are set.
      sent_id: Integer ID to use for the first sentence in this document.
      prefix: Optional string prepended to each printed sentence ID.

  Returns:
      The next sentence ID after all sentences in ``doc`` have been printed.

  The row layout follows the Universal Dependencies CoNLL-U format:
  https://universaldependencies.org/format.html
  """
  for sent in doc.sents:
    print("# sent_id = %s"%(prefix+str(sent_id)))
    print("# text = %s"%sent.sent)

    for i, word in enumerate(sent):
      #Find head
      if word.dep_.lower().strip() == 'root':
          head_idx = 0
      else:
          head_idx = word.head.i + 1 - sent[0].i

      #Find feature tag (if available)
      ftidx = word.tag_.find("__") + 2
      feature_tag=word.tag_[ftidx:]

      linetuple = (
          i+1,                                        #ID: Word index.
          word,                                       #FORM: Word form or punctuation symbol.
          word.lemma_.lower(),                        #LEMMA: Lemma or stem of word form.
          word.pos_,                                  #UPOSTAG: Universal part-of-speech tag drawn
                                                      # from revised version of the Google universal
                                                      # POS tags.
          '_',                                        #XPOSTAG: Language-specific part-of-speech tag;                                            # underscore if not available.
          '_' if feature_tag == "" else feature_tag,  #FEATS: List of morphological features from the
                                                      # universal feature inventory or from a defined
                                                      # language-specific extension; underscore if not
                                                      # available.
          head_idx,                                   #HEAD: Head of the current token, which is
                                                      # either a value of ID or zero (0).
          word.dep_.lower(),                          #DEPREL: Universal Stanford dependency relation
                                                      # to the HEAD (root iff HEAD = 0) or a defined
                                                      # language-specific subtype of one.
          '_',                                        #DEPS: List of secondary dependencies.
          '_'                                         #MISC: Any other annotation.
      )
      print("%i\t%s\t%s\t%s\t%s\t%s\t%i\t%s\t%s\t%s"%linetuple)

    sent_id+=1
    print("\n")

  return sent_id

def main(input_folder, out_dir_txt, conllu_output_folder, unit_type, docx_output_folder, model, input_format, prefix = ""):
    """Run text extraction and spaCy-to-CoNLL-U conversion for the input tree.

    Loads the requested spaCy model. DOCX inputs are first extracted to a
    mirrored text-file tree; TXT inputs are parsed directly. Each text line is
    then serialized to a matching CoNLL-U file. Segmentation output is written
    separately by :func:`segment_txt_as_docx`.

    Args:
        input_folder: Root directory containing the source file tree.
        out_dir_txt: Root directory for extracted UTF-8 text files.
        conllu_output_folder: Root directory for generated CoNLL-U text files.
        unit_type: Requested segment type, used by the subsequent segmentation
            stage.
        docx_output_folder: Destination root for the subsequent DOCX
            segmentation stage.
        model: Installed spaCy model name to load for linguistic annotation.
        input_format: Source file format, either ``"docx"`` or ``"txt"``.
        prefix: Optional prefix for emitted sentence IDs.

    Returns:
        None. Intermediate text and CoNLL-U files are written to disk.
    """

    print("Starting work in provided", str(input_folder), 'folder.')

    nlp = spacy.load(model)

    if input_format == "docx":
      print("1. Converting DOCX files to TXT files and saving to:", str(out_dir_txt))
      convert_docx_to_txt(input_folder, out_dir_txt)
      txt_paths = sorted(Path(out_dir_txt).rglob("*.txt"))
    else:
      print("1. Reading TXT input files from:", str(input_folder))
      intermediate_dirs = {
        Path(out_dir_txt).resolve(),
        Path(conllu_output_folder).resolve(),
      }
      txt_paths = [
        path for path in sorted(Path(input_folder).rglob("*.txt"))
        if not any(path.resolve().is_relative_to(directory) for directory in intermediate_dirs)
      ]

    CONLLU_folder_base = Path(conllu_output_folder)  # CONLLU folder base

    print("2. Converting input text to CONLL-U files and saving to:", str(conllu_output_folder))
    for path_txt in txt_paths:
      txt_root = Path(out_dir_txt) if input_format == "docx" else Path(input_folder)
      relative_path = path_txt.relative_to(txt_root)
      path_conllu = CONLLU_folder_base / relative_path
      path_conllu.parent.mkdir(parents=True, exist_ok=True)

      with path_txt.open("r", encoding='utf-8') as text_file, path_conllu.open("w", encoding="utf-8") as output_file:
        with redirect_stdout(output_file):
          sent_id = 1
          for line in text_file:
            doc = nlp(line.strip())
            sent_id = sentences_to_conllu(doc, sent_id, prefix=prefix)

if __name__ == "__main__":

  parser = argparse.ArgumentParser()
  parser.add_argument("--input_folder", required=True, type=Path, help="Path to a folder of input files; searched recursively.")
  parser.add_argument(
    "--input_format",
    required=True,
    choices=("docx", "txt"),
    help="Format of files in the input folder."
  )
  parser.add_argument(
    "--output_format",
    default="docx",
    choices=("docx", "txt"),
    help="Format for segmented output files. Defaults to docx."
  )
  parser.add_argument(
    "--include_conllu",
    action="store_true",
    help="Also copy CoNLL-U text files into Segmented_Documents/conllu-txts."
  )
  parser.add_argument(
    "--unit",
    default=None,
    required=True,
    choices=("s", "t", "c", "all"),
    help="`s`==sentences; `t`==t-units; `c`==clauses; `all`==all three."
  )
  parser.add_argument("--model", default="en_core_web_sm", required=False, type=str, help="Spacy model to use (e.g. 'es_core_news_md'). Default=='en_core_web_sm'")

  args = parser.parse_args()
  unit_type = args.unit

  in_dir = args.input_folder.expanduser().resolve()
  if not in_dir.is_dir():
    parser.error(f"Input folder does not exist or is not a directory: {in_dir}")

  out_dir_txt = Path(in_dir / "txts")
  out_dir_conllu = Path(in_dir / "conllu-txts")
  if args.input_format == "docx":
    out_dir_txt.mkdir(parents=True, exist_ok=True)
  out_dir_conllu.mkdir(parents=True, exist_ok=True)

  unit_type_foldername = {
    's': "sentences",
    't': "t_units",
    'c': "clauses",
    'all': "",
  }[unit_type]

  segmented_root = in_dir.parent / "Segmented_Documents"
  docx_output_folder = segmented_root / unit_type_foldername

  main(
    in_dir,
    out_dir_txt,
    out_dir_conllu,
    unit_type,
    docx_output_folder,
    args.model,
    args.input_format,
  )

  print(
    "3. Retrieving unit type",
    unit_type,
    "from CoNLL-U files and saving",
    args.output_format.upper(),
    "files at:",
    str(docx_output_folder),
  )
  segment_txt_as_docx(out_dir_conllu, docx_output_folder, unit_type, args.output_format)
  if args.include_conllu:
    conllu_output_copy_folder = segmented_root / "conllu-txts"
    print("4. Copying CoNLL-U files to:", str(conllu_output_copy_folder))
    copy_conllu_files(out_dir_conllu, conllu_output_copy_folder)

  print("5. Removing intermediate files from:", str(in_dir))
  for intermediate_dir in (out_dir_txt, out_dir_conllu):
    if intermediate_dir.exists():
      shutil.rmtree(intermediate_dir)
